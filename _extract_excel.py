# -*- coding: utf-8 -*-
import openpyxl
from pathlib import Path

xlsx_path = r"c:\Users\Javelin\Downloads\ホーチキ本番反映.xlsx"
out_path = r"C:\xampp\htdocs\hochiki\_excel_scan.txt"


def is_gray_fill(fill):
    if fill is None or fill.fill_type is None:
        return False
    if fill.patternType is None and fill.fill_type is None:
        return False
    fg = fill.fgColor
    if fg is None:
        return False
    if fg.type == "rgb" and fg.rgb:
        rgb = fg.rgb
        if isinstance(rgb, str) and len(rgb) >= 6:
            if len(rgb) == 8:
                rgb = rgb[2:]
            try:
                r = int(rgb[0:2], 16)
                g = int(rgb[2:4], 16)
                b = int(rgb[4:6], 16)
                if abs(r - g) < 30 and abs(g - b) < 30 and abs(r - b) < 30:
                    if r < 240:
                        return True
            except ValueError:
                pass
    if fg.type == "indexed" and fg.indexed is not None:
        gray_indices = {22, 23, 24, 25, 26, 27, 28, 29, 30, 31, 55}
        if fg.indexed in gray_indices:
            return True
    if fg.type == "theme" and fg.theme is not None:
        if fg.theme in (0, 1, 2, 3):
            return True
    return False


def cell_value(cell):
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).replace("\t", " ").replace("\n", " ").strip()


wb = openpyxl.load_workbook(xlsx_path, data_only=True)
lines = []
lines.append("=" * 80)
lines.append(f"EXCEL FILE: {xlsx_path}")
lines.append(f"SHEETS: {wb.sheetnames}")
lines.append("=" * 80)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    lines.append("")
    lines.append("#" * 80)
    lines.append(f"SHEET: {sheet_name}")
    lines.append(f"Dimensions: {ws.dimensions}")
    lines.append("#" * 80)

    gray_cells = []
    status_notes = []
    all_rows = []

    max_row = ws.max_row or 0
    max_col = ws.max_column or 0

    for row_idx in range(1, max_row + 1):
        row_vals = []
        row_has_content = False
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell_value(cell)
            if val:
                row_has_content = True
            row_vals.append(val)

            if is_gray_fill(cell.fill):
                coord = f"{openpyxl.utils.get_column_letter(col_idx)}{row_idx}"
                gray_cells.append(f"  {coord}: '{val}' (gray fill)")

            val_lower = val.lower()
            if val in ("済", "完了", "done", "Done", "DONE", "○", "〇", "✓", "済み"):
                status_notes.append(
                    f"  {openpyxl.utils.get_column_letter(col_idx)}{row_idx}: DONE status = '{val}'"
                )
            elif val in ("未", "未対応", "pending", "Pending", "PENDING", "×", "✗", "未着手", "要対応"):
                status_notes.append(
                    f"  {openpyxl.utils.get_column_letter(col_idx)}{row_idx}: PENDING status = '{val}'"
                )
            elif val_lower in ("済", "完了") or "済" in val or "未" in val:
                if any(k in val for k in ("済", "未", "完了", "対応")):
                    status_notes.append(
                        f"  {openpyxl.utils.get_column_letter(col_idx)}{row_idx}: STATUS-like = '{val}'"
                    )

        if row_has_content:
            all_rows.append((row_idx, "\t".join(row_vals)))

    lines.append(f"\n--- ROWS WITH CONTENT ({len(all_rows)} rows) ---")
    for row_idx, row_text in all_rows:
        lines.append(f"R{row_idx}:\t{row_text}")

    if gray_cells:
        lines.append(f"\n--- GRAY-FILLED CELLS ({len(gray_cells)}) ---")
        lines.extend(gray_cells)
    else:
        lines.append("\n--- GRAY-FILLED CELLS: none detected ---")

    if status_notes:
        lines.append(f"\n--- STATUS COLUMNS/VALUES ({len(status_notes)}) ---")
        lines.extend(status_notes)
    else:
        lines.append("\n--- STATUS COLUMNS/VALUES: none detected ---")

wb2 = openpyxl.load_workbook(xlsx_path, data_only=False)
lines.append("\n" + "=" * 80)
lines.append("FORMULA CELLS (if any)")
lines.append("=" * 80)
for sheet_name in wb2.sheetnames:
    ws = wb2[sheet_name]
    formulas = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                formulas.append(f"  {sheet_name}!{cell.coordinate}: {cell.value}")
    if formulas:
        lines.append(f"\n{sheet_name}:")
        lines.extend(formulas)

text = "\n".join(lines)
Path(out_path).write_text(text, encoding="utf-8")
print(f"Written {len(text)} chars to {out_path}")
print(f"Sheets: {wb.sheetnames}")
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"  {sn}: {ws.max_row} rows x {ws.max_column} cols")
